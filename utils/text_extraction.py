"""Text extraction and input validation for syllabi.

Supports PDF (via pdfplumber) and DOCX (via python-docx), plus pasted text.

PRIVACY: everything here works on bytes held in memory. Nothing is written
to disk and the syllabus content is never logged or printed.

Every problem raises ExtractionError with a friendly, user-facing message.
The UI shows that message instead of a stack trace.
"""
import io
import re

import pdfplumber
from docx import Document

# ---- Limits (change here if needed) ----
MAX_FILE_BYTES = 5 * 1024 * 1024   # 5 MB upload limit
MIN_WORDS = 30                     # shorter than this is not a usable syllabus
MAX_CHARS = 200_000                # about 30,000 words; a syllabus is normally far shorter
MIN_CHARS_PER_PAGE = 50            # less text per page => probably a scanned PDF
SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class ExtractionError(ValueError):
    """Raised with a friendly message the UI can show directly to the user."""


def clean_text(text: str) -> str:
    """Normalise line breaks and whitespace.

    We do NOT lowercase here, because the matcher needs the original case
    for ambiguous short skills like "R" or "Go" (matched case-sensitively).
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ")             # non-breaking spaces
    text = re.sub(r"[ \t\f\v]+", " ", text)          # collapse runs of spaces/tabs
    text = re.sub(r" *\n *", "\n", text)             # trim spaces around line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)           # max one blank line in a row
    return text.strip()


def validate_text(text: str) -> str:
    """Clean the text and make sure it is long enough to analyse.

    Used both for pasted text and for text extracted from files.
    Returns the cleaned text, or raises ExtractionError.
    """
    cleaned = clean_text(text or "")
    if not cleaned:
        raise ExtractionError(
            "The syllabus text is empty. Please upload a file, paste text, "
            "or choose a sample syllabus."
        )
    word_count = len(cleaned.split())
    if word_count < MIN_WORDS:
        raise ExtractionError(
            f"The syllabus text is too short ({word_count} words). Please provide "
            f"at least {MIN_WORDS} words, for example the course outline, "
            "learning outcomes and weekly topics."
        )
    if len(cleaned) > MAX_CHARS:
        raise ExtractionError(
            f"The syllabus text is too long ({len(cleaned):,} characters; the limit is "
            f"{MAX_CHARS:,}). Please use the course outline only, without appendices or readings."
        )
    return cleaned


def _extract_pdf(data: bytes) -> str:
    """Read all text from a PDF held in memory."""
    if not data.startswith(b"%PDF-"):
        raise ExtractionError(
            "This file does not look like a valid PDF. Please check the file "
            "or save it again as PDF."
        )
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_texts = [(page.extract_text() or "") for page in pdf.pages]
    except Exception:
        # Do not show technical details to the user.
        raise ExtractionError(
            "We could not read this PDF. It may be damaged or password-protected. "
            "Try saving it again, or paste the text instead."
        ) from None

    if not page_texts:
        raise ExtractionError("This PDF has no pages.")

    text = "\n".join(page_texts)
    # Scanned PDFs are images of pages, so they contain almost no text.
    if len(text.strip()) < MIN_CHARS_PER_PAGE * len(page_texts):
        raise ExtractionError(
            "This PDF seems to be scanned (an image of the pages), so we cannot "
            "read its text. Please use a text-based PDF or DOCX, or paste the text."
        )
    return text


def _extract_docx(data: bytes) -> str:
    """Read all text from a DOCX held in memory, including tables."""
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise ExtractionError(
            "We could not read this Word file. It may be damaged, or it may be "
            "an old .doc file. Please save it as .docx, or paste the text instead."
        ) from None

    parts = [p.text for p in doc.paragraphs]
    # Syllabi often keep weekly topics in tables, so read those too.
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Extract and validate syllabus text from an uploaded file.

    filename: original file name (used only to check the extension).
    data:     the file contents as bytes (kept in memory only).
    Returns cleaned text, or raises ExtractionError with a friendly message.
    """
    if not data:
        raise ExtractionError("The uploaded file is empty. Please choose another file.")

    if len(data) > MAX_FILE_BYTES:
        size_mb = len(data) / (1024 * 1024)
        raise ExtractionError(
            f"The file is too large ({size_mb:.1f} MB). The limit is 5 MB. "
            "Try removing images, or paste the text instead."
        )

    name = (filename or "").lower()
    if name.endswith(".pdf"):
        raw = _extract_pdf(data)
    elif name.endswith(".docx"):
        raw = _extract_docx(data)
    else:
        raise ExtractionError(
            "Unsupported file type. Please upload a PDF or DOCX file, "
            "or paste the text instead."
        )

    return validate_text(raw)
