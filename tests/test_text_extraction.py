"""Tests for utils/text_extraction.py.

All test files are built in memory, so no PDF/DOCX files are stored in the repo.
Run with:  python -m pytest -q
"""
import io

import pytest
from docx import Document

from utils.text_extraction import (
    MAX_CHARS,
    MAX_FILE_BYTES,
    ExtractionError,
    clean_text,
    extract_text,
    validate_text,
)

# A fictional syllabus snippet with more than 30 words.
SAMPLE = (
    "Course Outline. This course introduces data analysis using Python and SQL. "
    "Students will clean datasets, build dashboards in Tableau, apply descriptive "
    "statistics, and present findings to stakeholders in a final group project."
)


# ---------- helpers to build test files in memory ----------

def make_docx(paragraphs, table_rows=None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_pdf(lines) -> bytes:
    """Build a tiny one-page text PDF by hand (no extra library needed).
    With lines=[] it produces a page with no text, like a scanned PDF."""
    shown = " ".join(f"({line}) Tj T*" for line in lines)
    content = f"BT /F1 11 Tf 50 750 Td 14 TL {shown} ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref_start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_start}\n%%EOF\n"
    ).encode()
    return out


# ---------- clean_text / validate_text ----------

def test_clean_text_normalises_whitespace():
    assert clean_text("  Python\t\tand   SQL \r\n\r\n\r\n\r\nWeek 2  ") == "Python and SQL\n\nWeek 2"


def test_clean_text_keeps_case():
    # Case matters for ambiguous skills like "R", so it must not be lowercased.
    assert "R" in clean_text("Statistics with R")


def test_validate_text_accepts_normal_syllabus():
    assert validate_text(SAMPLE).startswith("Course Outline")


def test_validate_text_rejects_empty():
    with pytest.raises(ExtractionError, match="empty"):
        validate_text("   \n  ")


def test_validate_text_rejects_too_short():
    with pytest.raises(ExtractionError, match="too short"):
        validate_text("Python and SQL basics.")


def test_validate_text_rejects_too_long():
    with pytest.raises(ExtractionError, match="too long"):
        validate_text("word " * (MAX_CHARS // 4))


# ---------- DOCX ----------

def test_docx_extraction_reads_paragraphs_and_tables():
    data = make_docx([SAMPLE], table_rows=[["Week 1", "Introduction to Power BI"]])
    text = extract_text("syllabus.docx", data)
    assert "Tableau" in text
    assert "Power BI" in text  # text inside a table is read too


def test_docx_corrupted_file_is_rejected():
    with pytest.raises(ExtractionError, match="could not read this Word file"):
        extract_text("broken.docx", b"PK\x03\x04 not really a docx")


# ---------- PDF ----------

def test_pdf_extraction_reads_text():
    lines = [
        "Course Outline: Introduction to Data Analytics",
        "Students learn Python, SQL and descriptive statistics.",
        "Week 3 covers data cleaning and Tableau dashboards.",
        "Assessment includes a group project and a final presentation.",
    ]
    text = extract_text("syllabus.pdf", make_pdf(lines))
    assert "Python" in text
    assert "Tableau" in text


def test_pdf_without_text_is_treated_as_scanned():
    with pytest.raises(ExtractionError, match="scanned"):
        extract_text("scan.pdf", make_pdf([]))


def test_fake_pdf_is_rejected():
    with pytest.raises(ExtractionError, match="valid PDF"):
        extract_text("fake.pdf", b"this is not a pdf")


# ---------- general file checks ----------

def test_empty_file_is_rejected():
    with pytest.raises(ExtractionError, match="empty"):
        extract_text("syllabus.pdf", b"")


def test_file_over_5mb_is_rejected():
    with pytest.raises(ExtractionError, match="too large"):
        extract_text("big.pdf", b"%PDF-" + b"0" * MAX_FILE_BYTES)


def test_unsupported_type_is_rejected():
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_text("notes.txt", SAMPLE.encode())
